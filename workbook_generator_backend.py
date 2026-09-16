from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageOps, ImageFilter, ImageDraw
from pathlib import Path
import numpy as np
import math, subprocess, shutil, tempfile

INK='22262A'; MUTED='666C73'; PALE='F3F4F5'; LINE='D6DADD'

class WorkbookBuilder:
    def __init__(self, asset_dir):
        self.asset_dir = Path(asset_dir)
        self.asset_dir.mkdir(parents=True, exist_ok=True)
        self.cache = {}

    # ---------- image processing ----------
    def _content_crop_once(self, im, threshold=12, margin_frac=0.055):
        a=np.array(im.convert('L'))
        h,w=a.shape
        band=max(3,min(24,int(min(w,h)*0.025)))
        border=np.concatenate([a[:band,:].ravel(),a[-band:,:].ravel(),a[:, :band].ravel(),a[:,-band:].ravel()])
        bg=int(np.median(border))
        diff=np.abs(a.astype(np.int16)-bg)
        mask=diff>threshold
        row_counts=mask.sum(axis=1); col_counts=mask.sum(axis=0)
        rows=np.where(row_counts>max(2,int(w*0.002)))[0]
        cols=np.where(col_counts>max(2,int(h*0.002)))[0]
        if len(rows)==0 or len(cols)==0:
            return im
        y0,y1=rows[0],rows[-1]; x0,x1=cols[0],cols[-1]
        cw=x1-x0+1; ch=y1-y0+1
        mx=max(8,int(cw*margin_frac)); my=max(8,int(ch*margin_frac))
        x0=max(0,x0-mx); x1=min(w-1,x1+mx); y0=max(0,y0-my); y1=min(h-1,y1+my)
        if (x1-x0+1)>0.985*w and (y1-y0+1)>0.985*h:
            return im
        return im.crop((x0,y0,x1+1,y1+1))

    def load_crop(self, image_path):
        image_path = str(image_path)
        ck=('crop',image_path)
        if ck in self.cache:
            return self.cache[ck].copy()
        im=Image.open(image_path).convert('L')
        im=self._content_crop_once(im,12,0.055)
        im2=self._content_crop_once(im,14,0.075)
        if im2.width*im2.height >= 0.35*im.width*im.height:
            im=im2
        long=max(im.size)
        if long<1800:
            scale=1800/long
            im=im.resize((max(1,int(im.width*scale)),max(1,int(im.height*scale))),Image.Resampling.LANCZOS)
        self.cache[ck]=im.copy()
        return im

    def image_ratio(self, image_path):
        im=self.load_crop(image_path)
        return im.width/im.height

    def grid_dims_for_ratio(self, ratio, level):
        if level=='cross':
            return None
        nlong={'fine':8,'medium':5,'coarse':3}[level]
        if ratio>=1:
            cols=nlong; rows=max(2,round(nlong/ratio))
        else:
            rows=nlong; cols=max(2,round(nlong*ratio))
        return int(cols),int(rows)

    def overlay_scaffold(self, im, level, line=188):
        out=im.convert('L').copy(); d=ImageDraw.Draw(out)
        w,h=out.size
        d.rectangle([1,1,w-2,h-2],outline=138,width=3)
        if level is None or level=='none':
            return out
        if level=='cross':
            d.line([(w//2,0),(w//2,h)],fill=line,width=3)
            d.line([(0,h//2),(w,h//2)],fill=line,width=3)
            return out
        cols,rows=self.grid_dims_for_ratio(w/h,level)
        for i in range(1,cols):
            x=round(w*i/cols); d.line([(x,0),(x,h)],fill=line,width=3)
        for j in range(1,rows):
            y=round(h*j/rows); d.line([(0,y),(w,y)],fill=line,width=3)
        return out

    def blank_scaffold(self, ratio, level, long_px=1800):
        if ratio>=1:
            w=long_px; h=max(300,round(long_px/ratio))
        else:
            h=long_px; w=max(300,round(long_px*ratio))
        return self.overlay_scaffold(Image.new('L',(w,h),255),level)

    def quantize_values(self, im, n):
        x=ImageOps.autocontrast(im.convert('L'), cutoff=1)
        a=np.array(x)
        vals=a.reshape(-1)
        cuts=np.percentile(vals, np.linspace(0,100,n+1)[1:-1])
        idx=np.digitize(a,cuts)
        tones=np.linspace(30,235,n).astype(np.uint8)
        return Image.fromarray(tones[idx],mode='L')

    def progressive(self, im, stage):
        x=ImageOps.autocontrast(im.convert('L'),cutoff=1)
        w,h=x.size; long=max(w,h)
        if stage==1:
            target=34
            s=target/long
            sm=x.resize((max(8,int(w*s)),max(8,int(h*s))),Image.Resampling.BILINEAR)
            up=sm.resize((w,h),Image.Resampling.BICUBIC)
            return up.filter(ImageFilter.GaussianBlur(max(4,long/180)))
        if stage==2:
            target=110
            s=target/long
            sm=x.resize((max(12,int(w*s)),max(12,int(h*s))),Image.Resampling.BILINEAR)
            up=sm.resize((w,h),Image.Resampling.BICUBIC)
            return up.filter(ImageFilter.GaussianBlur(max(2,long/550)))
        if stage==3:
            target=420
            s=min(1,target/long)
            sm=x.resize((max(20,int(w*s)),max(20,int(h*s))),Image.Resampling.LANCZOS)
            up=sm.resize((w,h),Image.Resampling.BICUBIC)
            return up.filter(ImageFilter.GaussianBlur(max(0.7,long/1800)))
        return x

    def asset(self, image_path, key, kind='clear', level=None, upside=False):
        image_path = Path(image_path)
        ck=(str(image_path),key,kind,level,upside)
        if ck in self.cache:
            return self.cache[ck]
        im=self.load_crop(image_path)
        if upside:
            im=im.rotate(180)
        if kind=='clear':
            pass
        elif kind=='3value':
            im=self.quantize_values(im,3)
        elif kind=='5value':
            im=self.quantize_values(im,5)
        elif kind.startswith('focus'):
            im=self.progressive(im,int(kind[-1]))
        else:
            raise ValueError(kind)
        im=self.overlay_scaffold(im,level)
        p=self.asset_dir/f"{key}_{kind}_{level or 'none'}{'_up' if upside else ''}.png"
        if not p.exists():
            im.save(p,optimize=True)
        self.cache[ck]=p
        return p

    def blank_asset(self, image_path, key, level=None):
        ck=('blank',str(image_path),key,level)
        if ck in self.cache:
            return self.cache[ck]
        r=self.image_ratio(image_path)
        im=self.blank_scaffold(r,level)
        p=self.asset_dir/f"blank_{key}_{level or 'none'}.png"
        if not p.exists():
            im.save(p,optimize=True)
        self.cache[ck]=p
        return p

    # ---------- DOCX helpers ----------
    def set_cell_border(self, cell, **kwargs):
        tcPr=cell._tc.get_or_add_tcPr(); tcBorders=tcPr.first_child_found_in('w:tcBorders')
        if tcBorders is None:
            tcBorders=OxmlElement('w:tcBorders'); tcPr.append(tcBorders)
        for edge in ('top','left','bottom','right','insideH','insideV'):
            if edge in kwargs:
                el=tcBorders.find(qn('w:'+edge))
                if el is None:
                    el=OxmlElement('w:'+edge); tcBorders.append(el)
                for k,v in kwargs[edge].items():
                    el.set(qn('w:'+k),str(v))

    def set_cell_margins(self, cell, top=0,start=0,bottom=0,end=0):
        tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
        if tcMar is None:
            tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
        for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
            node=tcMar.find(qn('w:'+m))
            if node is None:
                node=OxmlElement('w:'+m); tcMar.append(node)
            node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

    def shade_cell(self, cell, fill):
        tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
        if shd is None:
            shd=OxmlElement('w:shd'); tcPr.append(shd)
        shd.set(qn('w:fill'),fill)

    def setup_doc(self):
        doc=Document(); sec=doc.sections[0]
        sec.page_width=Inches(8.5); sec.page_height=Inches(11)
        sec.top_margin=Inches(0.22); sec.bottom_margin=Inches(0.25); sec.left_margin=Inches(0.26); sec.right_margin=Inches(0.26)
        doc.styles['Normal'].font.name='Aptos'; doc.styles['Normal'].font.size=Pt(9.2); doc.styles['Normal'].font.color.rgb=RGBColor.from_string(INK)
        footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        r=footer.add_run('Print at 100% / Actual size'); r.font.name='Aptos'; r.font.size=Pt(7.2); r.font.color.rgb=RGBColor.from_string(MUTED)
        return doc

    def ptext(self, doc,text,size=10,bold=False,color=INK,space_after=4,align=None,font='Aptos'):
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(space_after)
        if align is not None:
            p.alignment=align
        r=p.add_run(text); r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=RGBColor.from_string(color)
        return p

    def page_header(self, doc,subject,title,subtitle=None, compact=False):
        p=self.ptext(doc,subject.upper(),7.8,True,MUTED,0)
        p.paragraph_format.keep_with_next=True
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.paragraph_format.keep_with_next=True
        r=p.add_run(title); r.font.name='Georgia'; r.font.size=Pt(16.5 if compact else 18); r.font.color.rgb=RGBColor.from_string(INK)
        if subtitle:
            p=self.ptext(doc,subtitle,8.5,False,MUTED,3); p.paragraph_format.keep_with_next=True

    def task_box(self, doc,text):
        t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; t.columns[0].width=Inches(7.82)
        c=t.cell(0,0); self.shade_cell(c,PALE); self.set_cell_margins(c,40,90,40,90)
        self.set_cell_border(c,top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'})
        p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
        rr=p.add_run(text); rr.font.name='Aptos'; rr.font.size=Pt(8.1); rr.font.color.rgb=RGBColor.from_string(MUTED)
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1)

    def page_break(self, doc):
        doc.add_page_break()

    def fit_inches(self, ratio,max_w,max_h):
        if ratio>=1:
            w=max_w; h=w/ratio
            if h>max_h:
                h=max_h; w=h*ratio
        else:
            h=max_h; w=h*ratio
            if w>max_w:
                w=max_w; h=w/ratio
        return w,h

    def add_picture_center(self, cell_or_doc, path, w,h, is_cell=False):
        if is_cell:
            p=cell_or_doc.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
        else:
            p=cell_or_doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
        p.add_run().add_picture(str(path),width=Inches(w),height=Inches(h))
        return p

    def dual_observation_page(self, doc, image_path, key, subject, include_blind=True, include_regular=True):
        title = 'Blind contour + regular contour' if include_blind and include_regular else ('Blind contour' if include_blind else 'Regular contour')
        subtitle = 'Two observation exercises on one page. No grid.' if include_blind and include_regular else 'Observation exercise. No grid.'
        self.page_header(doc, subject, title, subtitle, compact=True)
        blocks = []
        if include_blind:
            blocks.append(('Blind contour', 'Look at the reference, not the paper. Move slowly around the contour. Accuracy is not the goal.'))
        if include_regular:
            blocks.append(('Regular contour', 'Observe the outside contour and major inner edges. No grid: this is an observation exercise, not a placement test.'))
        ratio=self.image_ratio(image_path)
        max_w,max_h=(3.18,2.95) if ratio<=1.18 else (7.55,1.35)
        w,h=self.fit_inches(ratio,max_w,max_h)
        for bi,(title,task) in enumerate(blocks):
            self.ptext(doc,title,11.8,False,INK,1,font='Georgia')
            self.task_box(doc,task)
            if ratio<=1.18:
                t=doc.add_table(rows=2,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
                for i in range(2):
                    t.columns[i].width=Inches(3.78)
                    c=t.cell(0,i); self.set_cell_margins(c,0,0,12,0); self.set_cell_border(c,top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'})
                    p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    r=p.add_run('REFERENCE' if i==0 else 'YOUR DRAWING'); r.font.size=Pt(7.2); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(MUTED)
                for i,pth in enumerate([self.asset(image_path,key,'clear',None,False),self.blank_asset(image_path,key,None)]):
                    c=t.cell(1,i); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; self.set_cell_margins(c,0,0,0,0); self.set_cell_border(c,top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'})
                    self.add_picture_center(c,pth,w,h,True)
            else:
                for lab,pth in [('REFERENCE',self.asset(image_path,key,'clear',None,False)),('YOUR DRAWING',self.blank_asset(image_path,key,None))]:
                    self.ptext(doc,lab,7.1,True,MUTED,1,WD_ALIGN_PARAGRAPH.CENTER)
                    self.add_picture_center(doc,pth,w,h)
            spacer=doc.add_paragraph(); spacer.paragraph_format.space_after=Pt(2 if bi==0 else 0)
        self.page_break(doc)

    def _compact_exercise_block(self, doc, image_path, key, title, task, kind='clear', level=None, upside=False):
        self.ptext(doc,title,12.8,False,INK,1,font='Georgia')
        self.task_box(doc,task)
        ratio=self.image_ratio(image_path)
        if ratio<=1.18:
            max_w,max_h=3.55,3.85
            w,h=self.fit_inches(ratio,max_w,max_h)
            t=doc.add_table(rows=2,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
            for i in range(2):
                t.columns[i].width=Inches(3.88)
                c=t.cell(0,i); self.set_cell_margins(c,0,0,20,0); self.set_cell_border(c,top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'})
                p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                r=p.add_run('REFERENCE' if i==0 else 'YOUR DRAWING'); r.font.size=Pt(7.5); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(MUTED)
            for i,pth in enumerate([self.asset(image_path,key,kind,level,upside),self.blank_asset(image_path,key,level)]):
                c=t.cell(1,i); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; self.set_cell_margins(c,0,0,0,0); self.set_cell_border(c,top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'})
                self.add_picture_center(c,pth,w,h,True)
        else:
            max_w,max_h=7.7,1.7
            w,h=self.fit_inches(ratio,max_w,max_h)
            for lab,pth in [('REFERENCE',self.asset(image_path,key,kind,level,upside)),('YOUR DRAWING',self.blank_asset(image_path,key,level))]:
                self.ptext(doc,lab,7.2,True,MUTED,1,WD_ALIGN_PARAGRAPH.CENTER)
                self.add_picture_center(doc,pth,w,h)
        p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)

    def compact_pair_page(self, doc,image_path,key,subject,title,task,kind='clear',level=None,upside=False):
        self.page_header(doc,subject,title,compact=True)
        self.task_box(doc,task)
        ratio=self.image_ratio(image_path)
        if ratio<=1.18:
            max_w,max_h=3.56,5.9
            w,h=self.fit_inches(ratio,max_w,max_h)
            t=doc.add_table(rows=2,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
            for i in range(2):
                t.columns[i].width=Inches(3.88)
                c=t.cell(0,i); self.set_cell_margins(c,0,0,30,0); self.set_cell_border(c,top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'})
                p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                r=p.add_run('REFERENCE' if i==0 else 'YOUR DRAWING'); r.font.size=Pt(7.7); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(MUTED)
            for i,pth in enumerate([self.asset(image_path,key,kind,level,upside),self.blank_asset(image_path,key,level)]):
                c=t.cell(1,i); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP; self.set_cell_margins(c,0,0,0,0); self.set_cell_border(c,top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'})
                self.add_picture_center(c,pth,w,h,True)
        else:
            max_w,max_h=7.72,2.7
            w,h=self.fit_inches(ratio,max_w,max_h)
            for lab,pth in [('REFERENCE',self.asset(image_path,key,kind,level,upside)),('YOUR DRAWING',self.blank_asset(image_path,key,level))]:
                self.ptext(doc,lab,7.2,True,MUTED,1,WD_ALIGN_PARAGRAPH.CENTER)
                self.add_picture_center(doc,pth,w,h)
                p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
        self.page_break(doc)

    def sampler_page(self, doc,image_path,key,subject,skill,task,kind='clear',upside=False,levels=('fine','medium','coarse','cross')):
        self.page_header(doc,subject,f'{skill} - scaffold sampler','Four quick versions on one page. The cross is another support level, not a test.',compact=True)
        self.task_box(doc,task)
        ratio=self.image_ratio(image_path)
        table=doc.add_table(rows=2,cols=2); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
        for c in range(2): table.columns[c].width=Inches(3.9)
        for idx,lev in enumerate(levels):
            cell=table.cell(idx//2,idx%2); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            self.set_cell_margins(cell,45,40,45,40)
            self.set_cell_border(cell,top={'val':'single','sz':'5','color':LINE},bottom={'val':'single','sz':'5','color':LINE},left={'val':'single','sz':'5','color':LINE},right={'val':'single','sz':'5','color':LINE})
            p=cell.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
            r=p.add_run(lev.upper()); r.font.name='Aptos'; r.font.size=Pt(8); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(MUTED)
            inner=cell.add_table(rows=2,cols=2); inner.alignment=WD_TABLE_ALIGNMENT.CENTER; inner.autofit=False
            for j in range(2): inner.columns[j].width=Inches(1.68)
            for j,lab in enumerate(['REF','DRAW']):
                cc=inner.cell(0,j); self.set_cell_margins(cc,0,0,15,0); self.set_cell_border(cc,top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'})
                pp=cc.paragraphs[0]; pp.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=pp.add_run(lab); rr.font.size=Pt(6.6); rr.font.bold=True; rr.font.color.rgb=RGBColor.from_string(MUTED)
            iw,ih=self.fit_inches(ratio,1.56,2.4)
            for j,pth in enumerate([self.asset(image_path,key,kind,lev,upside),self.blank_asset(image_path,key,lev)]):
                cc=inner.cell(1,j); self.set_cell_margins(cc,0,0,0,0); self.set_cell_border(cc,top={'val':'nil'},bottom={'val':'nil'},left={'val':'nil'},right={'val':'nil'})
                self.add_picture_center(cc,pth,iw,ih,True)
        self.page_break(doc)

    def large_pair(self, doc,image_path,key,subject,title,task,kind='clear',level=None,upside=False):
        ratio=self.image_ratio(image_path)
        w,h=self.fit_inches(ratio,7.72,8.35)
        for lab,pth in [('REFERENCE',self.asset(image_path,key,kind,level,upside)),('DRAWING',self.blank_asset(image_path,key,level))]:
            self.page_header(doc,subject,title,subtitle=f'Large study - {lab.lower()}', compact=True)
            self.task_box(doc,task)
            self.ptext(doc,lab,7.8,True,MUTED,2,WD_ALIGN_PARAGRAPH.CENTER)
            self.add_picture_center(doc,pth,w,h)
            self.page_break(doc)

    def value_sequence(self, doc,image_path,key,subject, include_compact=True, include_large=True, do3=True, do5=True):
        if do3:
            if include_compact:
                self.compact_pair_page(doc,image_path,key,subject,'3-value study - compact','Reduce the image to three value families only: dark, middle, light. Ignore detail.',kind='3value',level='medium')
            if include_large:
                self.large_pair(doc,image_path,key,subject,'3-value study - large','Use the same three value families at a larger scale. Keep the medium grid for placement.',kind='3value',level='medium')
        if do5:
            if include_compact:
                self.compact_pair_page(doc,image_path,key,subject,'5-value study - compact','Use five value families from darkest to lightest. Keep the big masses unified.',kind='5value',level='medium')
            if include_large:
                self.large_pair(doc,image_path,key,subject,'5-value study - large','Repeat at a larger scale. Add intermediate values only after the large masses read.',kind='5value',level='medium')

    def progressive_focus(self, doc,image_path,key,subject):
        ratio=self.image_ratio(image_path); w,h=self.fit_inches(ratio,7.72,8.35)
        self.page_header(doc,subject,'Progressive focus - one evolving drawing','Use this same sheet through all four reference stages.', compact=True)
        self.task_box(doc,'Do not restart. Begin with the largest masses, then add only what each sharper reference genuinely reveals. The coarse grid stays so placement is not the challenge.')
        self.ptext(doc,'WORKING DRAWING',7.8,True,MUTED,2,WD_ALIGN_PARAGRAPH.CENTER)
        self.add_picture_center(doc,self.blank_asset(image_path,key,'coarse'),w,h); self.page_break(doc)
        stage_names={1:'Very blurred',2:'Medium blur',3:'Nearly sharp',4:'Sharp'}
        stage_tasks={1:'Place only the biggest masses and overall direction.',2:'Add major divisions, overlaps, and broad edge changes.',3:'Refine principal contours and secondary relationships.',4:'Add selected detail only where it helps the drawing.'}
        for st in range(1,5):
            self.page_header(doc,subject,f'Progressive focus - {stage_names[st]}',subtitle=f'Stage {st} of 4 - continue on the same drawing', compact=True)
            self.task_box(doc,stage_tasks[st])
            self.ptext(doc,'REFERENCE',7.8,True,MUTED,2,WD_ALIGN_PARAGRAPH.CENTER)
            self.add_picture_center(doc,self.asset(image_path,key,f'focus{st}','coarse'),w,h)
            self.page_break(doc)

    def final_integrated(self, doc,image_path,key,subject):
        self.compact_pair_page(doc,image_path,key,subject,'Integrated study - center cross','Use the center cross as a light placement scaffold. Apply whichever habits helped: envelope, landmarks, negative space, then values.',kind='clear',level='cross')
        self.large_pair(doc,image_path,key,subject,'Integrated study - center cross','Repeat larger with only the center cross. This is still supported; it is not a blank-page test.',kind='clear',level='cross')

    def subject_core(self, doc, spec):
        image_path=spec['image_path']; key=spec['key']; subject=spec['subject']
        exercises=spec.get('exercises', {})
        include_compact=spec.get('include_compact', True)
        include_large=spec.get('include_large', True)
        levels=spec.get('levels', ['fine','medium','coarse','cross'])

        if exercises.get('blind_contour') or exercises.get('regular_contour'):
            self.dual_observation_page(doc,image_path,key,subject, include_blind=exercises.get('blind_contour', False), include_regular=exercises.get('regular_contour', False))

        if exercises.get('sam'):
            if include_compact:
                self.sampler_page(doc,image_path,key,subject,'SAM / mapping','Locate top, bottom, left, right, then major intersections and angles before drawing. Use the grid deliberately.',kind='clear')
            if include_large:
                for lev in [lv for lv in ['fine','medium','coarse','cross'] if lv in levels]:
                    self.large_pair(doc,image_path,key,subject,f'SAM / mapping - {lev} scaffold','Map the envelope and landmarks first. Use only as much of the scaffold as is helpful.',kind='clear',level=lev)

        if exercises.get('upside_down'):
            if include_compact:
                self.sampler_page(doc,image_path,key,subject,'Upside-down drawing','Treat the image as shapes, angles, and distances rather than naming the object. Keep placement support.',kind='clear',upside=True)
            if include_large:
                for lev in [lv for lv in ['fine','medium','coarse'] if lv in levels]:
                    self.large_pair(doc,image_path,key,subject,f'Upside down - {lev} scaffold','Copy the upside-down image as shapes and distances. The scaffold is intentionally retained.',kind='clear',level=lev,upside=True)

        if exercises.get('negative_space'):
            if include_compact:
                self.sampler_page(doc,image_path,key,subject,'Negative space','Draw the empty shapes around and between the object parts. Let the object appear from those spaces.',kind='clear')
            if include_large:
                for lev in [lv for lv in ['fine','medium','coarse'] if lv in levels]:
                    self.large_pair(doc,image_path,key,subject,f'Negative space - {lev} scaffold','Use the scaffold to place the empty shapes. Changing what you observe should not also remove placement support.',kind='clear',level=lev)

        if exercises.get('three_value') or exercises.get('five_value'):
            self.value_sequence(doc,image_path,key,subject, include_compact=include_compact, include_large=include_large, do3=exercises.get('three_value', False), do5=exercises.get('five_value', False))

        if exercises.get('progressive_focus'):
            self.progressive_focus(doc,image_path,key,subject)

        if exercises.get('integrated'):
            self.final_integrated(doc,image_path,key,subject)

    def build(self, workbook_title, specs, output_docx, output_pdf=None):
        doc=self.setup_doc()
        for i,spec in enumerate(specs):
            self.subject_core(doc,spec)
        output_docx = Path(output_docx)
        doc.save(output_docx)
        if output_pdf:
            convert_docx_to_pdf(output_docx, output_pdf)
        return output_docx


def _find_libreoffice():
    # Works on Linux/macOS when LibreOffice is on PATH, and checks common
    # Windows install locations so the user usually does not need to edit PATH.
    import shutil as _shutil
    candidates = [
        _shutil.which('libreoffice'),
        _shutil.which('soffice'),
        r'C:\\Program Files\\LibreOffice\\program\\soffice.exe',
        r'C:\\Program Files (x86)\\LibreOffice\\program\\soffice.exe',
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            return str(candidate)
    raise RuntimeError('LibreOffice was not found. Install LibreOffice, then try again.')


def convert_docx_to_pdf(docx_path, output_pdf=None):
    docx_path = Path(docx_path)
    outdir = docx_path.parent if output_pdf is None else Path(output_pdf).parent
    outdir.mkdir(parents=True, exist_ok=True)
    cmd=[_find_libreoffice(),'--headless','--convert-to','pdf','--outdir',str(outdir),str(docx_path)]
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    produced = outdir / (docx_path.stem + '.pdf')
    if output_pdf is not None:
        output_pdf = Path(output_pdf)
        if produced.resolve() != output_pdf.resolve():
            shutil.copy2(produced, output_pdf)
            return output_pdf
    return produced
